import io
import json
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import load_workbook

from .models import Buchung, CrewMember, ShoppingItem, Toern, toern_slugify


class BordkasseTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('skipper', password='pw')
        self.toern = Toern.objects.create(name='Kroatien 2026', slug='kroatien-2026')
        self.anna = CrewMember.objects.create(toern=self.toern, name='Anna')
        self.ben = CrewMember.objects.create(toern=self.toern, name='Ben')

    def post(self, path, data=None):
        return self.client.post(f'/bordkasse/kroatien-2026/api/{path}', json.dumps(data or {}),
                                content_type='application/json')


class SlugTests(TestCase):
    def test_umlauts(self):
        self.assertEqual(toern_slugify('Ägäis Süd 2026'), 'aegaeis-sued-2026')
        self.assertEqual(toern_slugify('  !!  '), '')


class AccessTests(BordkasseTestCase):
    def test_anonymous_can_read(self):
        self.assertEqual(self.client.get('/bordkasse/').status_code, 200)
        r = self.client.get('/bordkasse/kroatien-2026/')
        self.assertContains(r, 'Kroatien 2026')
        self.assertEqual(self.client.get('/bordkasse/kroatien-2026/api/state/').status_code, 200)
        self.assertEqual(self.client.get('/bordkasse/kroatien-2026/export.xlsx').status_code, 200)

    def test_anonymous_cannot_change_kasse(self):
        self.assertEqual(self.post('crew/', {'name': 'Carla'}).status_code, 403)
        self.assertEqual(self.post(f'crew/{self.anna.id}/', {'name': 'X'}).status_code, 403)
        self.assertEqual(self.post('tx/', {'kind': 'einzahlung', 'person_id': self.anna.id,
                                           'amount': 10}).status_code, 403)
        self.assertFalse(Buchung.objects.exists())
        self.assertEqual(CrewMember.objects.count(), 2)

    def test_anonymous_cannot_create_toern(self):
        r = self.client.post('/bordkasse/', {'name': 'Elba'})
        self.assertEqual(r.status_code, 302)
        self.assertIn('/accounts/login/', r['Location'])
        self.assertFalse(Toern.objects.filter(name='Elba').exists())

    def test_anonymous_can_edit_shopping_list(self):
        r = self.post('shop/', {'text': 'Kaffee'})
        self.assertEqual(r.status_code, 200)
        item = ShoppingItem.objects.get()
        self.assertEqual(self.post(f'shop/{item.id}/', {'done': True, 'qty': '2 Pakete'}).status_code, 200)
        item.refresh_from_db()
        self.assertTrue(item.done)
        self.assertEqual(item.qty, '2 Pakete')
        self.assertEqual(self.post(f'shop/{item.id}/delete/').status_code, 200)
        self.assertFalse(ShoppingItem.objects.exists())

    def test_duplicate_open_shopping_item_rejected(self):
        self.post('shop/', {'text': 'Kaffee'})
        self.assertEqual(self.post('shop/', {'text': 'kaffee'}).status_code, 409)

    def test_logged_in_creates_toern(self):
        self.client.force_login(self.user)
        r = self.client.post('/bordkasse/', {'name': 'Ägäis Süd'})
        self.assertRedirects(r, '/bordkasse/aegaeis-sued/')
        self.assertEqual(Toern.objects.get(slug='aegaeis-sued').created_by, self.user)
        # gleicher Name nochmal → kein Duplikat, Weiterleitung auf bestehenden Törn
        r = self.client.post('/bordkasse/', {'name': 'ägäis süd'})
        self.assertRedirects(r, '/bordkasse/aegaeis-sued/')
        self.assertEqual(Toern.objects.filter(slug='aegaeis-sued').count(), 1)

    def test_toerns_are_separate(self):
        other = Toern.objects.create(name='Elba', slug='elba')
        self.client.force_login(self.user)
        r = self.client.post('/bordkasse/elba/api/tx/', json.dumps(
            {'kind': 'einzahlung', 'person_id': self.anna.id, 'amount': 5}), content_type='application/json')
        self.assertEqual(r.status_code, 400)  # Anna gehört nicht zu Elba
        self.assertFalse(other.buchungen.exists())


class KasseTests(BordkasseTestCase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.user)

    def state(self):
        return self.client.get('/bordkasse/kroatien-2026/api/state/').json()

    def test_totals_and_saldo(self):
        self.post('tx/', {'kind': 'einzahlung', 'person_id': self.anna.id, 'amount': '100'})
        self.post('tx/', {'kind': 'ausgabe', 'person_id': 'kasse', 'amount': '30', 'method': 'karte'})
        self.post('tx/', {'kind': 'ausgabe', 'person_id': self.ben.id, 'amount': '50', 'method': 'karte'})
        self.post('tx/', {'kind': 'ausgabe', 'person_id': self.ben.id, 'amount': '20,50', 'method': 'bar'})
        s = self.state()
        t = s['totals']
        self.assertEqual(t['kassenstand'], 70.0)       # 100 eingezahlt − 30 aus der Kasse
        self.assertEqual(t['sum_ein_bar'], 100.0)
        self.assertEqual(t['sum_aus_bar'], 50.5)       # Kasse (immer bar) + Ben privat bar
        self.assertEqual(t['sum_ein_karte'], 50.0)     # Karte: eingezahlt == ausgegeben
        self.assertEqual(t['sum_aus_karte'], 50.0)
        self.assertEqual(t['fair_share'], 50.25)       # 100.50 / 2
        saldo = {c['name']: c['saldo'] for c in s['crew']}
        self.assertEqual(saldo, {'Anna': 49.75, 'Ben': 20.25})
        kasse_tx = [x for x in s['tx'] if x['person_id'] is None][0]
        self.assertEqual(kasse_tx['method'], 'bar')   # Kasse ist immer Bargeld

    def test_einzahlung_requires_person(self):
        r = self.post('tx/', {'kind': 'einzahlung', 'person_id': 'kasse', 'amount': 10})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.post('tx/', {'kind': 'ausgabe', 'person_id': 'kasse', 'amount': -3}).status_code, 400)

    def test_rename_carries_over_to_bookings(self):
        self.post('tx/', {'kind': 'einzahlung', 'person_id': self.anna.id, 'amount': 40})
        r = self.post(f'crew/{self.anna.id}/', {'name': 'Anna-Lena'})
        self.assertEqual(r.status_code, 200)
        s = r.json()
        self.assertEqual(s['tx'][0]['person'], 'Anna-Lena')
        self.assertEqual([c['name'] for c in s['crew']], ['Anna-Lena', 'Ben'])
        self.assertEqual(self.post(f'crew/{self.anna.id}/', {'name': 'ben'}).status_code, 400)

    def test_edit_keeps_revision_and_soft_delete(self):
        s = self.post('tx/', {'kind': 'ausgabe', 'person_id': self.ben.id, 'amount': 12, 'method': 'karte',
                              'note': 'Diesel'}).json()
        tx_id = s['tx'][0]['id']
        # unveränderte Speicherung erzeugt keine Revision
        self.post(f'tx/{tx_id}/', {'kind': 'ausgabe', 'person_id': self.ben.id, 'amount': 12,
                                   'method': 'karte', 'note': 'Diesel'})
        s = self.post(f'tx/{tx_id}/', {'kind': 'ausgabe', 'person_id': self.anna.id, 'amount': 15,
                                       'method': 'bar', 'note': 'Diesel Palma'}).json()
        tx = s['tx'][0]
        self.assertEqual(len(tx['revisions']), 1)
        self.assertEqual(tx['revisions'][0]['person'], 'Ben')
        self.assertEqual(tx['revisions'][0]['amount'], 12.0)
        self.assertEqual(tx['revisions'][0]['changed_by'], 'skipper')
        self.assertEqual(tx['amount'], 15.0)
        self.assertIsNotNone(tx['updated_at'])

        s = self.post(f'tx/{tx_id}/delete/').json()
        tx = s['tx'][0]
        self.assertTrue(tx['deleted'])
        self.assertEqual(tx['deleted_by'], 'skipper')
        self.assertEqual(s['totals']['sum_aus'], 0.0)
        self.assertTrue(Buchung.objects.filter(pk=tx_id).exists())  # Soft-Delete
        self.assertEqual(self.post(f'tx/{tx_id}/', {'kind': 'ausgabe', 'person_id': 'kasse',
                                                    'amount': 1}).status_code, 400)

    def test_excel_export(self):
        tx = self.post('tx/', {'kind': 'einzahlung', 'person_id': self.anna.id, 'amount': 100}).json()['tx'][0]
        self.post(f'tx/{tx["id"]}/', {'kind': 'einzahlung', 'person_id': self.anna.id, 'amount': 80})
        self.post('shop/', {'text': 'Bier'})
        r = self.client.get('/bordkasse/kroatien-2026/export.xlsx')
        self.assertIn('attachment', r['Content-Disposition'])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Übersicht', 'Crew-Saldo', 'Buchungen', 'Einkaufsliste', 'Änderungen'])
        self.assertEqual(wb['Übersicht']['B4'].value, 80)
        events = [(row[1], row[5]) for row in wb['Änderungen'].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(events, [('Angelegt', 100), ('Bearbeitet', 80)])
        self.assertEqual(wb['Einkaufsliste']['A2'].value, 'Bier')

    def test_html_is_escaped_in_initial_state(self):
        ShoppingItem.objects.create(toern=self.toern, text='</script><script>alert(1)</script>')
        r = self.client.get('/bordkasse/kroatien-2026/')
        self.assertNotContains(r, '</script><script>alert(1)')
        self.assertEqual(Decimal('0'), Decimal('0'))


class AbrechnungTests(BordkasseTestCase):
    def book(self, kind, person, amount, method=''):
        Buchung.objects.create(toern=self.toern, kind=kind, person=person, amount=Decimal(amount), method=method)

    def test_transfers_pay_out_kasse_first(self):
        from .services import KASSE_LABEL, build_settlement
        carla = CrewMember.objects.create(toern=self.toern, name='Carla')
        self.book('einzahlung', self.anna, '100')
        self.book('einzahlung', self.ben, '50')
        self.book('ausgabe', None, '60', 'bar')                 # aus der Kasse → Rest 90
        self.book('ausgabe', carla, '90', 'karte')
        # gesamt 150 → Anteil 50; Anna +50, Ben 0, Carla +40; Summe = Restgeld 90
        a = build_settlement(self.toern)
        saldi = {p['name']: p['saldo'] for p in a['personen']}
        self.assertEqual(saldi, {'Anna': Decimal('50'), 'Ben': Decimal('0'), 'Carla': Decimal('40')})
        self.assertEqual(sum(saldi.values()), a['kassenstand'])
        self.assertEqual([(t['von'], t['an'], t['betrag']) for t in a['transfers']],
                         [(KASSE_LABEL, 'Anna', Decimal('50')), (KASSE_LABEL, 'Carla', Decimal('40'))])

    def test_person_to_person_and_cent_rounding(self):
        from .services import build_settlement
        carla = CrewMember.objects.create(toern=self.toern, name='Carla')
        self.book('ausgabe', self.anna, '100', 'karte')
        Buchung.objects.create(toern=self.toern, kind='ausgabe', person=self.ben, amount=Decimal('999'),
                               method='karte', deleted=True)  # gelöscht → zählt nicht
        a = build_settlement(self.toern)
        self.assertEqual([p['anteil'] for p in a['personen']],
                         [Decimal('33.34'), Decimal('33.33'), Decimal('33.33')])
        self.assertEqual(a['rundungs_cent'], 1)
        self.assertEqual(sorted((t['von'], t['an'], t['betrag']) for t in a['transfers']),
                         [('Ben', 'Anna', Decimal('33.33')), ('Carla', 'Anna', Decimal('33.33'))])
        self.assertEqual(sum(p['saldo'] for p in a['personen']), Decimal('0'))

    def test_page_is_public(self):
        self.book('einzahlung', self.anna, '20')
        r = self.client.get('/bordkasse/kroatien-2026/abrechnung/')
        self.assertContains(r, 'Endabrechnung')
        self.assertContains(r, '+20,00 €')   # Anna: 20 eingezahlt, noch keine Ausgaben → Anteil 0
        self.assertContains(self.client.get('/bordkasse/kroatien-2026/'), '/bordkasse/kroatien-2026/abrechnung/')


class KonfigTests(BordkasseTestCase):
    def set_pflicht(self, value):
        from .models import BordkasseKonfig
        konfig = BordkasseKonfig.load()
        konfig.authentifizierung_erforderlich = value
        konfig.save()

    PAGES = ['/bordkasse/', '/bordkasse/kroatien-2026/', '/bordkasse/kroatien-2026/abrechnung/',
             '/bordkasse/kroatien-2026/export.xlsx']

    def test_default_is_open(self):
        from .models import BordkasseKonfig
        self.assertFalse(BordkasseKonfig.load().authentifizierung_erforderlich)
        self.assertEqual(BordkasseKonfig.objects.count(), 1)
        for page in self.PAGES:
            self.assertEqual(self.client.get(page).status_code, 200, page)

    def test_pflicht_requires_login_everywhere(self):
        self.set_pflicht(True)
        for page in self.PAGES:
            r = self.client.get(page)
            self.assertEqual(r.status_code, 302, page)
            self.assertIn('/accounts/login/', r['Location'])
        self.assertEqual(self.client.get('/bordkasse/kroatien-2026/api/state/').status_code, 403)
        self.assertEqual(self.post('shop/', {'text': 'Kaffee'}).status_code, 403)
        self.assertFalse(ShoppingItem.objects.exists())

        self.client.force_login(self.user)
        for page in self.PAGES:
            self.assertEqual(self.client.get(page).status_code, 200, page)
        self.assertEqual(self.post('shop/', {'text': 'Kaffee'}).status_code, 200)
