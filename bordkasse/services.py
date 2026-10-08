"""Berechnungen, Serialisierung und Excel-Export für die Bordkasse."""
import io
from decimal import Decimal
from zoneinfo import ZoneInfo

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from .models import Buchung

KASSE_LABEL = 'Bordkasse (Bargeld)'
LOCAL_TZ = ZoneInfo('Europe/Berlin')
ZERO = Decimal('0.00')

# Standard-Proviantliste für Segeltörns — reine Auswahlquelle; ein Klick legt
# einen ganz normalen Eintrag in der Einkaufsliste des Törns an.
STANDARD_LIST = [
    ('Frühstück', ['Brot/Brötchen', 'Butter', 'Margarine', 'Marmelade', 'Honig', 'Müsli', 'H-Milch',
                   'Eier', 'Kaffee', 'Tee', 'Orangensaft', 'Speck', 'Wurst', 'Käse']),
    ('Mittagessen', ['Nudeln', 'Tomatensauce', 'Pesto', 'Thunfisch (Dose)', 'Reis', 'Knäckebrot']),
    ('Abendessen', ['Hack', 'Hähnchenbrust', 'Dorade', 'Paprika', 'Zucchini', 'Tomaten', 'Pilze',
                    'Gurke', 'Salat', 'Kartoffeln', 'Zwiebeln', 'Knoblauch', 'Olivenöl', 'Öl',
                    'Essig', 'Salz & Pfeffer', 'Gewürzmischung']),
    ('Obst', ['Melonen', 'Äpfel', 'Bananen', 'Pfirsiche', 'Weintrauben']),
    ('Getränke', ['Wasser still', 'Wasser medium', 'Bier', 'Bier alkoholfrei', 'Weißwein', 'Rotwein',
                  'Prosecco', 'Aperol', 'Cola', 'Cola Zero', 'Limo', 'Limo Zero', 'Säfte']),
    ('Snacks', ['Chips', 'Salzgebäck', 'Nüsse', 'Kekse', 'Schokolade', 'Müsliriegel']),
    ('Sonstiges', ['Toilettenpapier', 'Küchenrolle', 'Spülmittel', 'Schwämme', 'Müllbeutel',
                   'Kaffeefilter']),
]


def can_write(user):
    """Törn anlegen und Crew/Buchungen ändern: nur eingeloggt."""
    return bool(user and user.is_authenticated)


def person_label(obj):
    """Anzeigename für Buchung/Revision: Crew-Name oder Bordkasse."""
    return obj.person.name if obj.person_id else KASSE_LABEL


def method_label(obj):
    if obj.kind != Buchung.AUSGABE:
        return ''
    if obj.person_id is None:
        return 'Bar (Bordkasse)'
    return 'Bar (privat)' if obj.method == Buchung.BAR else 'Karte'


def compute_totals(crew, buchungen):
    """Splitwise-Logik aus dem Prototyp.

    Saldo pro Person = (eingezahlt + selbst ausgelegt) − fairer Anteil an allen Ausgaben.
    Kassenstand = Bar-Einzahlungen − Ausgaben direkt aus der Bordkasse.
    Karte hat keinen eigenen Topf: was jemand per Karte vorstreckt, gilt sofort als
    eingebracht — daher ist bei Karte „eingezahlt“ immer gleich „ausgegeben“.
    """
    ein_by_person, aus_by_person = {}, {}
    sum_ein_bar = sum_aus_bar = sum_aus_karte = aus_kasse = ZERO
    for b in buchungen:
        if b.deleted:
            continue
        if b.kind == Buchung.EINZAHLUNG:
            ein_by_person[b.person_id] = ein_by_person.get(b.person_id, ZERO) + b.amount
            sum_ein_bar += b.amount  # Einzahlungen gehen immer als Bargeld in die Kasse
        else:
            if b.person_id is None:
                aus_kasse += b.amount
            else:
                aus_by_person[b.person_id] = aus_by_person.get(b.person_id, ZERO) + b.amount
            if b.method == Buchung.KARTE:
                sum_aus_karte += b.amount
            else:
                sum_aus_bar += b.amount

    sum_aus = sum_aus_bar + sum_aus_karte
    fair_share = (sum_aus / len(crew)) if crew else ZERO
    saldo = {}
    for c in crew:
        eingebracht = ein_by_person.get(c.id, ZERO) + aus_by_person.get(c.id, ZERO)
        saldo[c.id] = {
            'eingebracht': eingebracht,
            'saldo': eingebracht - fair_share,
        }
    return {
        'kassenstand': sum_ein_bar - aus_kasse,
        'sum_ein_bar': sum_ein_bar,
        'sum_aus_bar': sum_aus_bar,
        'sum_ein_karte': sum_aus_karte,
        'sum_aus_karte': sum_aus_karte,
        'sum_ein': sum_ein_bar + sum_aus_karte,
        'sum_aus': sum_aus,
        'fair_share': fair_share,
        'saldo': saldo,
    }


def load_toern_data(toern):
    crew = list(toern.crew.all())
    buchungen = list(
        toern.buchungen
        .select_related('person', 'created_by', 'deleted_by')
        .prefetch_related('revisionen__person', 'revisionen__changed_by')
    )
    shop = list(toern.einkaufsliste.all())
    return crew, buchungen, shop


def _money(d):
    return float(round(d, 2))


def _ts(dt):
    return dt.isoformat() if dt else None


def _user(u):
    return u.get_username() if u else ''


def _snapshot(obj):
    return {
        'kind': obj.kind,
        'person_id': obj.person_id,
        'person': person_label(obj),
        'amount': _money(obj.amount),
        'method': obj.method,
        'note': obj.note,
    }


def serialize_state(toern, user):
    crew, buchungen, shop = load_toern_data(toern)
    totals = compute_totals(crew, buchungen)
    return {
        'trip': {'name': toern.name, 'slug': toern.slug},
        'can_write': can_write(user),
        'totals': {k: _money(v) for k, v in totals.items() if k != 'saldo'},
        'crew': [{
            'id': c.id,
            'name': c.name,
            'eingebracht': _money(totals['saldo'][c.id]['eingebracht']),
            'saldo': _money(totals['saldo'][c.id]['saldo']),
        } for c in crew],
        'tx': [dict(_snapshot(b), **{
            'id': b.id,
            'created_at': _ts(b.created_at),
            'created_by': _user(b.created_by),
            'updated_at': _ts(b.updated_at),
            'deleted': b.deleted,
            'deleted_at': _ts(b.deleted_at),
            'deleted_by': _user(b.deleted_by),
            'revisions': [dict(_snapshot(r), changed_at=_ts(r.changed_at), changed_by=_user(r.changed_by))
                          for r in b.revisionen.all()],
        }) for b in buchungen],
        'shop': [{
            'id': s.id, 'text': s.text, 'qty': s.qty, 'info': s.info, 'done': s.done,
            'created_at': _ts(s.created_at),
        } for s in shop],
    }


# ---------------------------------------------------------------- Endabrechnung

def _cents(d):
    return int((d * 100).to_integral_value())


def _eur(cents):
    return Decimal(cents) / 100


def build_settlement(toern):
    """Endabrechnung, cent-genau.

    Der faire Anteil wird in Cent aufgeteilt; bleibt ein Rest, zahlen die ersten
    Crew-Mitglieder je 1 Cent mehr. Die Summe aller Salden ist genau das Restgeld in
    der Bordkasse — es wird zuerst ausgezahlt, danach gleichen die Personen untereinander aus.
    """
    crew, buchungen, _ = load_toern_data(toern)
    aktiv = sorted((b for b in buchungen if not b.deleted), key=lambda b: b.created_at)
    einzahlungen = [b for b in aktiv if b.kind == Buchung.EINZAHLUNG]
    ausgaben = [b for b in aktiv if b.kind == Buchung.AUSGABE]

    total_c = sum(_cents(b.amount) for b in ausgaben)
    kasse_ein_c = sum(_cents(b.amount) for b in einzahlungen)
    kasse_aus = [b for b in ausgaben if b.person_id is None]
    kasse_aus_c = sum(_cents(b.amount) for b in kasse_aus)
    kassenstand_c = kasse_ein_c - kasse_aus_c

    shares = {}
    rest = 0
    if crew:
        base, rest = divmod(total_c, len(crew))
        shares = {c.id: base + (1 if i < rest else 0) for i, c in enumerate(crew)}

    personen = []
    for c in crew:
        ein = [b for b in einzahlungen if b.person_id == c.id]
        aus = [b for b in ausgaben if b.person_id == c.id]
        ein_c = sum(_cents(b.amount) for b in ein)
        karte_c = sum(_cents(b.amount) for b in aus if b.method == Buchung.KARTE)
        bar_c = sum(_cents(b.amount) for b in aus if b.method != Buchung.KARTE)
        eingebracht_c = ein_c + karte_c + bar_c
        personen.append({
            'id': c.id,
            'name': c.name,
            'einzahlungen': ein,
            'auslagen': aus,
            'eingezahlt': _eur(ein_c),
            'karte': _eur(karte_c),
            'bar_privat': _eur(bar_c),
            'eingebracht': _eur(eingebracht_c),
            'anteil': _eur(shares[c.id]),
            'saldo_c': eingebracht_c - shares[c.id],
            'saldo': _eur(eingebracht_c - shares[c.id]),
        })

    # Ausgleich: Kasse zahlt ihr Restgeld zuerst aus, dann zahlen Schuldner an Gläubiger
    # (jeweils größter Betrag zuerst → wenige Überweisungen).
    payers = ([[KASSE_LABEL, kassenstand_c, True]] if kassenstand_c > 0 else []) + sorted(
        ([p['name'], -p['saldo_c'], False] for p in personen if p['saldo_c'] < 0), key=lambda x: -x[1])
    receivers = ([[KASSE_LABEL, -kassenstand_c]] if kassenstand_c < 0 else []) + sorted(
        ([p['name'], p['saldo_c']] for p in personen if p['saldo_c'] > 0), key=lambda x: -x[1])
    transfers = []
    i = j = 0
    while i < len(payers) and j < len(receivers):
        amount = min(payers[i][1], receivers[j][1])
        if amount > 0:
            transfers.append({'von': payers[i][0], 'an': receivers[j][0], 'betrag': _eur(amount),
                              'aus_kasse': payers[i][2]})
        payers[i][1] -= amount
        receivers[j][1] -= amount
        if payers[i][1] == 0:
            i += 1
        if receivers[j][1] == 0:
            j += 1

    return {
        'personen': personen,
        'transfers': transfers,
        'gesamt_ausgaben': _eur(total_c),
        'anzahl_ausgaben': len(ausgaben),
        'kasse_eingezahlt': _eur(kasse_ein_c),
        'kasse_ausgaben': kasse_aus,
        'kasse_ausgegeben': _eur(kasse_aus_c),
        'kassenstand': _eur(kassenstand_c),
        'anteil': _eur(total_c // len(crew)) if crew else ZERO,
        'rundungs_cent': rest,
        'crew_count': len(crew),
    }


# ---------------------------------------------------------------- Excel-Export

def _local(dt):
    return timezone.localtime(dt, LOCAL_TZ).replace(tzinfo=None) if dt else None


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord('A') + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if hasattr(cell.value, 'year'):
                cell.number_format = 'DD.MM.YYYY HH:MM'
            elif isinstance(cell.value, Decimal):
                cell.number_format = '#,##0.00 €'
    ws.freeze_panes = 'A2'
    return ws


def build_export(toern):
    """Excel mit den Reitern Übersicht, Crew-Saldo, Buchungen, Einkaufsliste, Änderungen."""
    crew, buchungen, shop = load_toern_data(toern)
    t = compute_totals(crew, buchungen)
    wb = Workbook()
    wb.remove(wb.active)

    _sheet(wb, 'Übersicht', ['Kennzahl', 'Wert'], [
        ['Törn', toern.name],
        ['Exportiert am', _local(timezone.now())],
        ['Kassenstand (Bargeld)', t['kassenstand']],
        ['Bar eingezahlt (€)', t['sum_ein_bar']],
        ['Bar ausgegeben (€)', t['sum_aus_bar']],
        ['Karte eingezahlt (€)', t['sum_ein_karte']],
        ['Karte ausgegeben (€)', t['sum_aus_karte']],
        ['Gesamt eingezahlt (€)', t['sum_ein']],
        ['Gesamt ausgegeben (€)', t['sum_aus']],
        ['Crew-Mitglieder', len(crew)],
    ], [26, 24])

    _sheet(wb, 'Crew-Saldo', ['Name', 'Eingezahlt + ausgelegt (€)', 'Fair-Anteil (€)', 'Saldo (€)'], [
        [c.name, t['saldo'][c.id]['eingebracht'], round(t['fair_share'], 2), round(t['saldo'][c.id]['saldo'], 2)]
        for c in crew
    ], [24, 26, 16, 14])

    aktiv = sorted((b for b in buchungen if not b.deleted), key=lambda b: b.created_at)
    _sheet(wb, 'Buchungen',
           ['Datum', 'Art', 'Person', 'Betrag (€)', 'Zahlart', 'Notiz', 'Bearbeitet?', 'Erfasst von'], [
               [_local(b.created_at), b.get_kind_display(), person_label(b), b.amount, method_label(b),
                b.note, 'Ja' if b.revisionen.all() else '', _user(b.created_by)]
               for b in aktiv
           ], [17, 12, 22, 12, 16, 30, 12, 16])

    _sheet(wb, 'Einkaufsliste', ['Artikel', 'Menge', 'Notiz', 'Status'], [
        [s.text, s.qty, s.info, 'Erledigt' if s.done else 'Offen']
        for s in sorted(shop, key=lambda s: s.created_at)
    ], [30, 16, 30, 10])

    events = []
    for b in buchungen:
        revs = list(b.revisionen.all())
        # „Angelegt“ zeigt den ursprünglichen Stand (= erste Revision, falls bearbeitet).
        first = revs[0] if revs else b
        events.append((b.created_at, 'Angelegt', first, _user(b.created_by)))
        # Jede Bearbeitung zeigt den *neuen* Stand: Stand der nächsten Revision bzw. aktueller Stand.
        for i, r in enumerate(revs):
            after = revs[i + 1] if i + 1 < len(revs) else b
            events.append((r.changed_at, 'Bearbeitet', after, _user(r.changed_by)))
        if b.deleted:
            events.append((b.deleted_at, 'Gelöscht', b, _user(b.deleted_by)))
    events.sort(key=lambda e: e[0])
    _sheet(wb, 'Änderungen',
           ['Datum', 'Ereignis', 'Buchung #', 'Art', 'Person', 'Betrag (€)', 'Zahlart', 'Notiz', 'Von'], [
               [_local(ts), ev, (b.buchung_id if hasattr(b, 'buchung_id') else b.id), b.get_kind_display(),
                person_label(b), b.amount, method_label(b), b.note, who]
               for ts, ev, b, who in events
           ], [17, 12, 10, 12, 22, 12, 16, 30, 16])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
